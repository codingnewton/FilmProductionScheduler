from openpyxl import load_workbook
import pandas as pd
from datetime import time
from scheduler import Optimizer, ROLE_NAME, Member, Task
from typing import Tuple, List
from .scheduler_classes import Member, Task, ROLE_NAME

def get_members_and_tasks(excel_file) -> Tuple[List[Task], List[Member]]:
    """
    Extract member and task data from an Excel file.
    
    Args:
        excel_file: Path to the Excel file containing member and task information
        
    Returns:
        Tuple of (tasks, members) lists
    """
    sheet_name = "Members_availability"
    workbook = load_workbook(filename=excel_file, data_only=True)
    sheet = workbook['Members']
    sheet2 = workbook['Tasks']

    member_names = []
    roles = []
    member_hr_rates = []
    member_ot_rates = []
    members = []
    tasks = []

    for cell in sheet[3][1:]:
        if cell.value:
            member_names.append(cell.value)

            col_index = cell.col_idx
            roles.append(sheet.cell(row=4, column=col_index).value)
            member_hr_rates.append(sheet.cell(row=5, column=col_index).value)
            member_ot_rates.append(sheet.cell(row=6, column=col_index).value)

    member_roles = [key for string in roles for key, value in ROLE_NAME.items() if value == string]

    def format_time(row_index):
        hours = row_index // 4
        minutes = (row_index % 4) * 15
        if hours == 24 and minutes == 0:
            return 23, 59
        else:
            return hours, minutes

    def process_unavailability(file_path, sheet_name="Members_availability"):
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        times = df.iloc[:, 0]
        results = {}
        for col_idx in range(1, df.shape[1]): 
            availability = df.iloc[:, col_idx].values 
            intervals = []
            start = None
            for i, available in enumerate(availability):
                if available == 0 and start is None:
                    start = i 
                elif available == 1 and start is not None:
                    intervals.append((format_time(start-1), format_time(i-1)))
                    start = None
            if start is not None:
                intervals.append((format_time(start-1), format_time(len(availability)-1)))
            results[f"Person_{col_idx}"] = intervals
        return results

    unavailability_results = process_unavailability(excel_file, sheet_name)
    member_intervals = []
    for person, intervals in unavailability_results.items():
        member_intervals.append(intervals)

    for i, name in enumerate(member_names):
        converted_intervals = [
            (time(start[0], start[1]), time(end[0], end[1])) for start, end in member_intervals[i]]
        members.append(Member(id=i, name=name, rate=member_hr_rates[i], ot=member_ot_rates[i], role=member_roles[i], blocked_timeslots=converted_intervals))

    task_names = []
    task_locations = []
    task_durations = []
    task_dependencies = []
    # task_timeofday = []
    task_actors = []
    task_roles = []

    for row in range(3, sheet2.max_row+1):
        cell_value = sheet2.cell(row=row, column=3).value
        if cell_value:
            task_names.append(cell_value)
            task_locations.append(sheet2.cell(row=row, column = 9).value)
            task_durations.append(sheet2.cell(row=row, column = 10).value)
            task_dependencies.append(sheet2.cell(row=row, column = 12).value)
            task_actors.append(sheet2.cell(row=row, column = 15).value)
            task_roles.append(sheet2.cell(row=row, column = 24).value)

    for i, name in enumerate(task_names):
        task_location = eval(task_locations[i])
        task_duration = task_durations[i]
        task_actor = eval(task_actors[i])
        task_role = eval(task_roles[i])
        task_dependency = eval(task_dependencies[i])
        tasks.append(Task(id=i, description=name, location=task_location, estimated_duration=task_duration, members=task_actor, roles=task_role, dependencies=task_dependency))
    
    return tasks, members

def generate_schedule(
    excel_file,
    output_location,
    timeout=30,
    num_workers=8, 
    break_duration=30, 
    max_time_before_break=240
):
    """
    Generate an optimized schedule based on the provided Excel file.
    
    Args:
        excel_file: Path to the Excel file containing member and task information
        output_location: Path to save the output files
        timeout: Optimization timeout in seconds
        num_workers: Number of workers for optimization
        break_duration: Break duration in minutes
        max_time_before_break: Maximum time before a break in minutes
        
    Returns:
        Dictionary containing optimization results
    """
    from .scheduler import Optimizer
    
    tasks, members = get_members_and_tasks(excel_file=excel_file)
    model = Optimizer(tasks, members)
    model.schedule_tasks(
        timeout=timeout, 
        num_workers=num_workers, 
        break_duration=break_duration, 
        max_time_before_break=max_time_before_break
    )
    model.export_schedule(output_location)

    result = {"cost": model.optimized_cost}
    return result